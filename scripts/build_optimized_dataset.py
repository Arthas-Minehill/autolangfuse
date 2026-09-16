"""
构建优化评测数据集脚本。

从 docs/update.xlsx 读取原始用例数据，结合当前受支持查询能力产生的已验证结果，
生成标准化的 Langfuse DatasetItem 格式 JSON。

用法：
    uv run python scripts/build_optimized_dataset.py

输出：docs/optimized_dataset.json

注意：查询结果需要预先通过当前受支持且受治理的数据能力获取并写入 sql_results_cache.json。
如果缓存文件不存在，脚本只会输出不含 expected_result 的骨架数据。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from aieval_runner.evaluation.tables.contract import normalize_table

XLSX_PATH = ROOT / "docs" / "update.xlsx"
OUTPUT_PATH = ROOT / "docs" / "optimized_dataset.json"
CACHE_PATH = ROOT / "docs" / "sql_results_cache.json"

DATASET_NAME = "aieval/data_analysis"


def _read_xlsx(path: Path) -> List[Dict[str, Any]]:
    """读取 xlsx 文件，返回行字典列表。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("需要 openpyxl: pip install openpyxl")

    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.active
    rows_iter = sheet.iter_rows(values_only=True)
    headers = [str(h or "").strip() for h in next(rows_iter)]
    records = []
    for row in rows_iter:
        record = {}
        for i, val in enumerate(row):
            if i < len(headers) and headers[i]:
                record[headers[i]] = val
        if any(v not in (None, "") for v in record.values()):
            records.append(record)
    wb.close()
    return records


def _parse_jsonish(value: Any, default: Any = None) -> Any:
    """安全解析 JSON 字符串。"""
    if value in (None, ""):
        return default
    if isinstance(value, (dict, list)):
        return value
    if isinstance(value, str):
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return value
    return value


def _extract_main_query(expected_tools_raw: Any) -> str:
    """从 expected_tools JSON 中提取最后一条 SQL 相关调用的指标名称。"""
    calls = _parse_jsonish(expected_tools_raw, [])
    if not isinstance(calls, list):
        return ""
    target_suffixes = {
        "getmetricappqueryresult",
        "searchmetricappqueryresult",
        "getsqlbymetrictypeandnameexact",
        "searchsqlbymetrictypeandnameexact",
    }
    for call in reversed(calls):
        if not isinstance(call, dict):
            continue
        tool_name = str(call.get("tool_name") or call.get("name") or "").lower()
        tool_key = tool_name.rsplit(".", 1)[-1] if "." in tool_name else tool_name
        if tool_key not in target_suffixes:
            continue
        args = call.get("arguments")
        if not isinstance(args, dict):
            continue
        req = args.get("req")
        if isinstance(req, dict) and isinstance(req.get("name"), str) and req["name"].strip():
            return req["name"].strip()
        name = args.get("name")
        if isinstance(name, str) and name.strip():
            return name.strip()
    return ""


def _canonical_tool_names(expected_tools_raw: Any) -> List[str]:
    """提取并规范化工具名称列表。"""
    calls = _parse_jsonish(expected_tools_raw, [])
    if not isinstance(calls, list):
        return []
    names = []
    seen = set()
    for call in calls:
        if not isinstance(call, dict):
            continue
        raw = str(call.get("tool_name") or call.get("name") or "").strip()
        if not raw:
            continue
        # 规范化 server 名称
        if raw.startswith("middle-mcp-remote."):
            raw = raw.replace("middle-mcp-remote.", "metric-mcp-remote.", 1)
        # 规范化 tool 名称
        tool_map = {
            "getmetricapppage": "searchMetricApp",
            "getmetricappqueryresult": "searchMetricAppQueryResult",
            "getsqlbymetrictypeandnameexact": "searchSqlByMetricTypeAndNameExact",
        }
        if "." in raw:
            server, tool = raw.rsplit(".", 1)
            tool_lower = tool.lower()
            if tool_lower in tool_map:
                raw = f"{server}.{tool_map[tool_lower]}"
        if raw not in seen:
            seen.add(raw)
            names.append(raw)
    return names


def _parse_schema_columns(schema_raw: Any) -> List[str]:
    """解析 expected_result_schema 中的 columns 列表。"""
    parsed = _parse_jsonish(schema_raw, {})
    if not isinstance(parsed, dict):
        return []
    columns = parsed.get("columns")
    if not isinstance(columns, list):
        return []
    return [str(c).strip() for c in columns if str(c).strip()]


def _build_evaluation_focus(
    schema_columns: List[str],
    answer_key_points: Any,
    grading_rubric: Any,
) -> List[str]:
    """构建 evaluation_focus 列表。"""
    focus = [f"field:{c}" for c in schema_columns]
    if isinstance(answer_key_points, str) and answer_key_points.strip():
        focus.append(f"answer_key_point:{answer_key_points.strip()}")
    if isinstance(grading_rubric, str) and grading_rubric.strip():
        focus.append(f"grading_rubric:{grading_rubric.strip()}")
    return focus


def build_items(xlsx_records: List[Dict[str, Any]], sql_cache: Dict[str, Any]) -> List[Dict[str, Any]]:
    """从 xlsx 记录构建优化的 DatasetItem 列表。"""
    items = []
    for row in xlsx_records:
        case_id = str(row.get("case_id") or "").strip()
        if not case_id:
            continue

        user_query = str(row.get("user_query") or "").strip()
        main_query = _extract_main_query(row.get("expected_tools"))
        expected_tools = _canonical_tool_names(row.get("expected_tools"))
        schema_columns = _parse_schema_columns(row.get("expected_result_schema"))
        source_sql = str(row.get("expected_sql") or "").strip()
        answer_key_points = row.get("expected_answer_key_points")
        grading_rubric = row.get("grading_rubric")
        evaluation_focus = _build_evaluation_focus(schema_columns, answer_key_points, grading_rubric)

        # 从缓存获取 SQL 执行结果
        cached_result = sql_cache.get(case_id) if sql_cache else None
        expected_result = None
        if cached_result:
            try:
                expected_result = normalize_table(
                    cached_result,
                    field_name=f"case_id={case_id}.expected_result",
                )
            except ValueError:
                expected_result = cached_result

        expected_output: Dict[str, Any] = {
            "case_id": case_id,
            "main_query": main_query,
            "expected_tools": expected_tools,
            "evaluation_focus": evaluation_focus,
        }
        if source_sql:
            expected_output["source_sql"] = source_sql
        if expected_result:
            expected_output["expected_result"] = expected_result

        metadata: Dict[str, Any] = {"case_id": case_id}
        for key in ("task_level", "task_type", "business_domain", "difficulty", "source"):
            val = row.get(key)
            if isinstance(val, str):
                val = val.strip()
            if val not in (None, ""):
                metadata[key] = val

        input_obj = {
            "case_id": case_id,
            "original_request": user_query,
        }

        items.append({
            "case_id": case_id,
            "input": input_obj,
            "expectedOutput": expected_output,
            "metadata": metadata,
        })

    return items


def main() -> None:
    if not XLSX_PATH.exists():
        print(f"错误: 找不到 {XLSX_PATH}")
        sys.exit(1)

    records = _read_xlsx(XLSX_PATH)
    print(f"从 xlsx 读取到 {len(records)} 条记录")

    sql_cache = {}
    if CACHE_PATH.exists():
        sql_cache = json.loads(CACHE_PATH.read_text(encoding="utf-8"))
        print(f"从缓存加载了 {len(sql_cache)} 条 SQL 执行结果")
    else:
        print("注意: sql_results_cache.json 不存在，expected_result 将为空")
        print("请先通过当前受支持的数据查询能力执行查询并将结果写入该缓存文件")

    items = build_items(records, sql_cache)

    output = {
        "dataset": DATASET_NAME,
        "description": "优化后的评测数据集 - 基于 docs/update.xlsx 提取，补充已验证查询结果作为 expected_result",
        "optimized_at": "2026-06-16",
        "schema_version": "aieval-external-runner-v2",
        "item_count": len(items),
        "items": items,
    }

    OUTPUT_PATH.write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"已写入 {OUTPUT_PATH} ({len(items)} 条 items)")


if __name__ == "__main__":
    main()
